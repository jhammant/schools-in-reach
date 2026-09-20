#!/usr/bin/env python3
"""Build a school register for each of Northern Ireland's 11 councils.

Usage (from the project root):
    python3 pipeline/uk/northern_ireland.py [--refresh] [--offline]

Outputs (site/data/):
    la/<ONS council code>/schools.json   register of grant-aided schools (same shape as England)
    uk/northern_ireland_las.json         council index (same fields as england/las.json entries)

Sources (all Open Government Licence v3.0):
  * Department of Education NI, "School enrolment - school level data 2025/2026"
    (school census, 10 October 2025): reference data, enrolments, sex, and the
    matching "Available places" workbooks for approved enrolment numbers.
  * OpenDataNI / Department of Education NI, "School Locations" (Locate a School,
    February 2016): building-level coordinates, used only where the published
    postcode still matches the current DE record.
  * postcodes.io (ONS Postcode Directory / OS Open Data / Royal Mail PAF open
    data): postcode centroids and ONS council codes.

Everything downloaded is cached under pipeline/.cache/uk/ni so re-runs are offline.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from hashlib import sha1
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency hint
    sys.exit("openpyxl is required: run `uv run --with openpyxl python3 pipeline/uk/northern_ireland.py`")

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "pipeline" / ".cache" / "uk" / "ni"
OUT = ROOT / "site" / "data"

USER_AGENT = "SchoolCatchmentArea-pipeline/1.0 (+personal open-data project)"
OGL = "Open Government Licence v3.0"

DE_PUBLICATION = "https://www.education-ni.gov.uk/publications/school-enrolment-school-level-data-20252026"
CENSUS_DATE = "2025-10-10"  # DE metadata: "In 2025/26, the school census date was 10th October."
CENSUS_YEAR = "2025/26"

# DE 2025/26 school-level workbooks (links taken from the publication page above).
DE_FILES = {
    "primary": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-07/School%20level%20-%20primary%20schools%20202526%20supp_1.xlsx",
        "de_primary_202526.xlsx",
    ),
    "postprimary": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-07/School%20level%20-%20post%20primary%20schools%20202526%20supp_1.xlsx",
        "de_postprimary_202526.xlsx",
    ),
    "nursery": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-07/School%20level%20-%20nursery%20schools%20202526%20supp_1.xlsx",
        "de_nursery_202526.xlsx",
    ),
    "special": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-03/School%20level%20-%20special%20schools%20202526.xlsx",
        "de_special_202526.xlsx",
    ),
}
DE_PLACES = {
    "primary": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-06/Available%20places%20-%20Primary%20202526%20-%20Revised%203%20June%202026.XLSX",
        "de_places_primary_202526.xlsx",
    ),
    "postprimary": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-03/Available%20places%20-%20Post-primary%20202526.XLSX",
        "de_places_postprimary_202526.xlsx",
    ),
    "nursery": (
        "https://www.education-ni.gov.uk/sites/default/files/2026-03/Available%20places%20-%20Nursery%20schools%20and%20units%20202526.XLSX",
        "de_places_nursery_202526.xlsx",
    ),
}
LOCATE_URL = (
    "https://admin.opendatani.gov.uk/dataset/39cd6af4-8fed-4ac9-9620-577a2190bb34/"
    "resource/d0947faf-5d84-4ce4-80dd-ce4fa0e1c0d5/download/locate-a-school-open-data-feb-2016.csv"
)
LOCATE_FILE = "locate_a_school_2016.csv"
LOCATE_PORTAL = "https://www.opendatani.gov.uk/@department-of-education/locate-a-school"

POSTCODES_API = "https://api.postcodes.io/postcodes"
POSTCODES_BATCH = 100
POSTCODES_PAUSE = 0.6  # seconds between bulk calls

# ONS council codes (2014 local government districts). DE publishes the council
# name only, so the code comes from this table and is checked against the
# admin_district code postcodes.io returns for every school postcode.
COUNCILS = {
    "antrim and newtownabbey": ("N09000001", "Antrim and Newtownabbey"),
    "armagh city, banbridge and craigavon": ("N09000002", "Armagh City, Banbridge and Craigavon"),
    "belfast": ("N09000003", "Belfast"),
    "causeway coast and glens": ("N09000004", "Causeway Coast and Glens"),
    "derry city and strabane": ("N09000005", "Derry City and Strabane"),
    "fermanagh and omagh": ("N09000006", "Fermanagh and Omagh"),
    "lisburn and castlereagh": ("N09000007", "Lisburn and Castlereagh"),
    "mid and east antrim": ("N09000008", "Mid and East Antrim"),
    "mid ulster": ("N09000009", "Mid Ulster"),
    "newry, mourne and down": ("N09000010", "Newry, Mourne and Down"),
    "ards and north down": ("N09000011", "Ards and North Down"),
}

# DE "School Type" -> the site's phase vocabulary.
PHASE = {
    "Primary": "Primary",
    "Prep": "Primary",
    "Secondary": "Secondary",
    "Grammar": "Grammar",
    "Nursery": "Nursery",
    "Special": "Special",
}
KIND = {
    "Primary": "primary school",
    "Prep": "preparatory department",
    "Secondary": "secondary (non-grammar) school",
    "Grammar": "grammar school",
    "Nursery": "nursery school",
    "Special": "special school",
}
# DE "Management Type" -> the NI management sector, spelled out.
MANAGEMENT = {
    "Controlled": "Controlled",
    "Controlled Integrated": "Controlled Integrated",
    "Catholic Maintained": "Catholic Maintained",
    "Other Maintained": "Other Maintained",
    "Gmi": "Grant Maintained Integrated",
    "Voluntary": "Voluntary",
}
SUPPRESSED = {"*", "#", "!", "", "n/a", "na", "-"}


# --------------------------------------------------------------------- helpers
def log(message: str) -> None:
    print(message, flush=True)


def fetch(url: str, name: str, refresh: bool = False, offline: bool = False) -> Path:
    """Download `url` into the cache once; later runs read the cached copy."""
    path = CACHE / name
    if path.exists() and not refresh:
        return path
    if offline:
        raise SystemExit(f"--offline but {path} is not cached yet")
    path.parent.mkdir(parents=True, exist_ok=True)
    log(f"  downloading {name}")
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=180) as response:
        payload = response.read()
    path.write_bytes(payload)
    time.sleep(1)
    return path


def clean(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def number(value) -> int | None:
    """DE suppresses small counts with * and #; those become null."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    text = str(value).strip().replace(",", "")
    if text.lower() in SUPPRESSED:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def sheet_table(path: Path, sheet: str, markers: tuple[str, ...]) -> tuple[list[str], list[tuple]]:
    """Return (header, rows) for a DE worksheet, skipping its preamble rows.

    `markers` are column headings that only ever appear on the header row, which
    is how the preamble (title, notes, blank rows) above each table is skipped.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = next((s for s in workbook.sheetnames if s.strip().lower() == sheet.lower()), None)
    if name is None:
        workbook.close()
        raise KeyError(f"{path.name}: no sheet called {sheet!r} (have {workbook.sheetnames})")
    rows = list(workbook[name].iter_rows(values_only=True))
    workbook.close()
    wanted = {h.lower().replace(" ", "") for h in markers}
    index = next(
        (
            i
            for i, row in enumerate(rows)
            if row and any((clean(cell) or "").lower().replace(" ", "") in wanted for cell in row)
        ),
        None,
    )
    if index is None:
        raise KeyError(f"{path.name}/{sheet}: no header row containing one of {markers}")
    header = [clean(cell) or "" for cell in rows[index]]
    body = [row for row in rows[index + 1 :] if row and clean(row[0])]
    return header, body


def column(header: list[str], *names: str) -> int | None:
    lookup = {h.lower().replace(" ", ""): i for i, h in enumerate(header)}
    for name in names:
        hit = lookup.get(name.lower().replace(" ", ""))
        if hit is not None:
            return hit
    return None


def keyed(path: Path, sheet: str, markers: tuple[str, ...], columns: dict[str, tuple[str, ...]]) -> dict[str, dict]:
    """{DE reference: {alias: value}} for the named columns of one worksheet."""
    header, rows = sheet_table(path, sheet, markers)
    picks = {alias: column(header, *names) for alias, names in columns.items()}
    missing = [alias for alias, index in picks.items() if index is None]
    if missing:
        raise KeyError(f"{path.name}/{sheet}: missing columns {missing} (header {header})")
    out: dict[str, dict] = {}
    for row in rows:
        ref = clean(row[0])
        if ref:
            out[ref] = {alias: row[index] for alias, index in picks.items()}
    return out


# ------------------------------------------------------------------- geocoding
def geocode(postcodes: list[str], refresh: bool, offline: bool) -> dict[str, dict]:
    """Bulk-lookup postcodes via postcodes.io, caching every raw response."""
    unique = sorted({p.upper() for p in postcodes if p})
    cache_dir = CACHE / "postcodes"
    cache_dir.mkdir(parents=True, exist_ok=True)
    results: dict[str, dict] = {}
    batches = [unique[i : i + POSTCODES_BATCH] for i in range(0, len(unique), POSTCODES_BATCH)]
    for n, batch in enumerate(batches, 1):
        digest = sha1("|".join(batch).encode()).hexdigest()[:16]
        path = cache_dir / f"bulk_{digest}.json"
        if not path.exists() or refresh:
            if offline:
                raise SystemExit(f"--offline but {path} is not cached yet")
            log(f"  postcodes.io batch {n}/{len(batches)} ({len(batch)} postcodes)")
            body = json.dumps({"postcodes": batch}).encode()
            req = urllib.request.Request(
                POSTCODES_API,
                data=body,
                headers={"Content-Type": "application/json", "User-Agent": USER_AGENT},
            )
            with urllib.request.urlopen(req, timeout=120) as response:
                path.write_bytes(response.read())
            time.sleep(POSTCODES_PAUSE)
        payload = json.loads(path.read_text(encoding="utf-8"))
        for entry in payload.get("result") or []:
            query = (entry.get("query") or "").upper()
            results[query] = entry.get("result") or {}

    # A handful of school postcodes have been terminated by Royal Mail but are
    # still the address DE publishes; postcodes.io keeps their centroid.
    for postcode in [p for p in unique if not results.get(p)]:
        path = cache_dir / f"terminated_{postcode.replace(' ', '_')}.json"
        if not path.exists() or refresh:
            if offline:
                raise SystemExit(f"--offline but {path} is not cached yet")
            log(f"  postcodes.io terminated lookup for {postcode}")
            url = f"https://api.postcodes.io/terminated_postcodes/{urllib.parse.quote(postcode)}"
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            try:
                with urllib.request.urlopen(req, timeout=60) as response:
                    path.write_bytes(response.read())
            except urllib.error.HTTPError as error:
                path.write_bytes(error.read() or b'{"status":404}')
            time.sleep(POSTCODES_PAUSE)
        found = (json.loads(path.read_text(encoding="utf-8")) or {}).get("result") or {}
        if found.get("latitude") is not None:
            found["_terminated"] = True
            results[postcode] = found
    return results


# ----------------------------------------------------------------- record build
def age_range(school_type: str, enrolments: dict | None) -> tuple[int | None, int | None]:
    """Age range from the year groups DE actually publishes for the school."""
    if school_type == "Nursery":
        return 3, 4
    if school_type == "Special":
        return 3, 19
    if school_type in ("Primary", "Prep"):
        nursery = 0
        if enrolments:
            nursery = (number(enrolments.get("nursery_ft")) or 0) + (number(enrolments.get("nursery_pt")) or 0)
        return (3 if nursery > 0 else 4), 11
    if school_type in ("Secondary", "Grammar"):
        sixth = 0
        if enrolments:
            sixth = (number(enrolments.get("year13")) or 0) + (number(enrolments.get("year14")) or 0)
        return 11, (18 if sixth > 0 else 16)
    return None, None


def tidy_phone(value) -> str | None:
    """Normalise the 2016 DE phone numbers to 028 xxxx xxxx where possible."""
    text = clean(value)
    if not text:
        return None
    digits = re.sub(r"\D", "", text)
    if digits.startswith("44"):
        digits = "0" + digits[2:]
    if len(digits) == 8:  # local number recorded without the 028 area code
        digits = "028" + digits
    if len(digits) == 11 and digits.startswith("028"):
        return f"028 {digits[3:7]} {digits[7:]}"
    return text


def distance_m(a: tuple[float, float], b: tuple[float, float]) -> float:
    """Rough metres between two WGS84 points (good enough for a sanity check)."""
    lat = (a[0] + b[0]) / 2
    dy = (a[0] - b[0]) * 111_320
    dx = (a[1] - b[1]) * 111_320 * max(0.1, abs(math.cos(math.radians(lat))))
    return (dx * dx + dy * dy) ** 0.5


def build_schools(refresh: bool, offline: bool) -> tuple[list[dict], dict]:
    log("Reading Department of Education NI school-level workbooks")
    reference_columns = {
        "name": ("School name",),
        "building": ("Building number",),
        "address": ("Address Line",),
        "town": ("Town",),
        "postcode": ("Postcode",),
        "school_type": ("School Type",),
        "management": ("Management Type",),
        "council": ("District Council (2014)",),
        "ward": ("Ward (2014)",),
        "dea": ("DEA (2014)",),
        "urban_rural": ("Urban/ Rural", "Urban/Rural"),
        "datazone": ("Datazone",),
    }
    files = {key: fetch(url, name, refresh, offline) for key, (url, name) in DE_FILES.items()}
    places_files = {key: fetch(url, name, refresh, offline) for key, (url, name) in DE_PLACES.items()}

    reference: dict[str, dict] = {}
    for key, path in files.items():
        rows = keyed(path, "Reference Data", ("DENI ref", "DENIref"), reference_columns)
        for ref, row in rows.items():
            row["_file"] = key
            reference[ref] = row
        log(f"  {path.name}: {len(rows)} schools")

    enrolments: dict[str, dict] = {}
    enrolments.update(
        keyed(
            files["primary"],
            "Enrolments",
            ("DENI ref", "DENIref"),
            {
                "total": ("Total enrolment",),
                "nursery_ft": ("Mainstream class: Nursery FT",),
                "nursery_pt": ("Mainstream class: Nursery PT",),
            },
        )
    )
    enrolments.update(
        keyed(
            files["postprimary"],
            "Enrolments",
            ("DENI ref", "DENIref"),
            {"total": ("Total enrolment",), "year13": ("Total: Year 13",), "year14": ("Total: Year 14",)},
        )
    )
    enrolments.update(keyed(files["nursery"], "Enrolments", ("DENI ref", "DENIref"), {"total": ("Total pupils",)}))
    enrolments.update(keyed(files["special"], "Sex", ("DENI ref", "DENIref"), {"total": ("Total enrolment",)}))

    sexes: dict[str, dict] = {}
    for key, boys, girls in (("primary", "Male", "Female"), ("nursery", "Boys", "Girls"), ("special", "Male", "Female")):
        sexes.update(keyed(files[key], "Sex", ("DENI ref", "DENIref"), {"boys": (boys,), "girls": (girls,)}))

    capacity: dict[str, int] = {}
    for key, path in places_files.items():
        header, rows = sheet_table(path, "School level data", ("Approved enrolments",))
        index = column(header, "Approved enrolments")
        for row in rows:
            ref = clean(row[0])
            value = number(row[index]) if index is not None else None
            if ref and value is not None and (key != "nursery" or ref in reference):
                # the nursery workbook covers nursery units inside primary schools too;
                # only let it fill a capacity the primary workbook did not already set
                if key == "nursery" and ref in capacity:
                    continue
                capacity[ref] = value

    log("Reading OpenDataNI 'School Locations' (Feb 2016) for building-level coordinates")
    locate: dict[str, dict] = {}
    with fetch(LOCATE_URL, LOCATE_FILE, refresh, offline).open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            ref = clean(row.get("Reference"))
            if ref:
                locate[ref] = row

    log("Geocoding postcodes with postcodes.io")
    lookups = geocode([clean(r["postcode"]) or "" for r in reference.values()], refresh, offline)

    stats = {
        "non_numeric_ref": [],
        "unplaced": [],
        "council_mismatch": [],
        "locate_points": 0,
        "postcode_centroids": 0,
        "unknown_council": [],
    }
    schools: list[dict] = []
    for ref, row in sorted(reference.items()):
        if not ref.isdigit():
            stats["non_numeric_ref"].append(ref)
            continue
        school_type = clean(row["school_type"]) or ""
        management = clean(row["management"]) or ""
        postcode = (clean(row["postcode"]) or "").upper()
        lookup = lookups.get(postcode) or {}

        council_name = clean(row["council"]) or ""
        council = COUNCILS.get(re.sub(r"\s+", " ", council_name).lower())
        if council is None:
            stats["unknown_council"].append((ref, council_name))
            continue
        la_code, la_name = council
        lookup_code = ((lookup.get("codes") or {}).get("admin_district") or "") if lookup else ""
        if lookup_code and lookup_code != la_code:
            stats["council_mismatch"].append((ref, clean(row["name"]), postcode, la_name, lookup.get("admin_district")))

        lat = lookup.get("latitude") if lookup else None
        lon = lookup.get("longitude") if lookup else None
        geo_source = None
        if lat is not None:
            geo_source = (
                "postcodes.io terminated-postcode centroid"
                if lookup.get("_terminated")
                else "postcodes.io postcode centroid"
            )
        if lat is not None:
            stats["postcode_centroids"] += 1
        # refine with the 2016 building point when its postcode still matches
        located = locate.get(ref)
        if located and (clean(located.get("Postcode")) or "").replace(" ", "").upper() == postcode.replace(" ", ""):
            try:
                point = (float(located["Latitude"]), float(located["Longitude"]))
            except (TypeError, ValueError):
                point = None
            if point and (lat is None or distance_m(point, (lat, lon)) < 1500):
                lat, lon = point
                geo_source = "OpenDataNI School Locations (2016) building point"
                stats["postcode_centroids"] -= 1
                stats["locate_points"] += 1
        if lat is None:
            stats["unplaced"].append((ref, clean(row["name"]), postcode))

        enrolment = enrolments.get(ref) or {}
        sex = sexes.get(ref) or {}
        boys, girls = number(sex.get("boys")), number(sex.get("girls"))
        pupils = number(enrolment.get("total"))
        gender = None
        if boys is not None and girls is not None and pupils and pupils >= 25:
            gender = "Boys" if girls == 0 else "Girls" if boys == 0 else "Mixed"
        low, high = age_range(school_type, enrolment)
        sector = MANAGEMENT.get(management, management)
        # `type` keeps DE's own management wording; `type_group` names the sector,
        # where "Voluntary" post-primary schools are the voluntary grammar sector.
        establishment_type = f"{sector} {KIND.get(school_type, (school_type or '').lower())}".strip()
        if sector == "Voluntary" and school_type == "Grammar":
            sector = "Voluntary Grammar"
        street = " ".join(x for x in (clean(row["building"]), clean(row["address"])) if x) or None
        sixth_form = None
        if school_type in ("Secondary", "Grammar"):
            sixth_form = ((number(enrolment.get("year13")) or 0) + (number(enrolment.get("year14")) or 0)) > 0
        nursery_provision = None
        if school_type in ("Primary", "Prep"):
            nursery_provision = (
                (number(enrolment.get("nursery_ft")) or 0) + (number(enrolment.get("nursery_pt")) or 0)
            ) > 0

        schools.append(
            {
                "urn": 30000000 + int(ref),
                "source_id": ref,
                "name": clean(row["name"]),
                "lat": round(lat, 6) if lat is not None else None,
                "lon": round(lon, 6) if lon is not None else None,
                "phase": PHASE.get(school_type, school_type or None),
                "type": establishment_type or None,
                "type_group": sector or None,
                "gender": gender,
                "religious_ethos": "Roman Catholic" if management == "Catholic Maintained" else None,
                "selective": True if school_type == "Grammar" else False if school_type == "Secondary" else None,
                "age_low": low,
                "age_high": high,
                "nursery_provision": nursery_provision,
                "sixth_form": sixth_form,
                "capacity": capacity.get(ref),
                "pupils": pupils,
                "boys": boys,
                "girls": girls,
                "census_date": CENSUS_DATE,
                "street": street,
                "locality": None,
                "town": clean(row["town"]),
                "postcode": postcode or None,
                "website": None,
                "phone": tidy_phone(located.get("Telephone")) if located else None,
                "head": None,
                "inspectorate": "Education and Training Inspectorate",
                "ward": clean(row["ward"]),
                "dea": clean(row["dea"]),
                "urban_rural": clean(row["urban_rural"]),
                "datazone": clean(row["datazone"]),
                "geo_source": geo_source,
                "la": la_code,
                "la_name": la_name,
                "nation": "Northern Ireland",
            }
        )
    return schools, stats


# ----------------------------------------------------------------------- output
def write_json(path: Path, payload) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    path.write_text(text, encoding="utf-8")
    return len(text.encode("utf-8"))


def sources() -> list[dict]:
    return [
        {
            "title": f"School enrolment - school level data {CENSUS_YEAR} (Department of Education, Northern Ireland)",
            "url": DE_PUBLICATION,
            "licence": OGL,
        },
        {
            "title": "School Locations / Locate a School, February 2016 (Department of Education NI, OpenDataNI)",
            "url": LOCATE_PORTAL,
            "licence": OGL,
        },
        {
            "title": "postcodes.io - ONS Postcode Directory, OS Open Data and Royal Mail open postcode data",
            "url": "https://postcodes.io/",
            "licence": OGL,
        },
    ]


def notes(count: int) -> list[str]:
    return [
        f"Grant-aided nursery, primary, post-primary and special schools from the Department of Education NI "
        f"school census of {CENSUS_DATE} ({CENSUS_YEAR}); {count} schools in this council.",
        "urn is 30000000 + the Department of Education reference number, so it cannot collide with a 6-digit "
        "English URN; source_id is the DE reference as published.",
        "type_group is the NI management sector as DE publishes it (Controlled, Controlled Integrated, Catholic "
        "Maintained, Other Maintained - predominantly the Irish-medium sector, Grant Maintained Integrated, "
        "Voluntary / Voluntary Grammar).",
        "selective is true where DE classes the school as Grammar, false for non-grammar post-primary schools, and "
        "null for primary, nursery, preparatory and special schools. It is DE's classification, not confirmation "
        "that a school used an entrance assessment in a given year: the Schools' Entrance Assessment Group (SEAG) "
        "assessment is used by 63 post-primary schools, against 66 DE-classified grammar schools here, and SEAG's "
        "membership list is not open-licensed so it has not been merged in.",
        "lat/lon are the building point from the 2016 OpenDataNI School Locations dataset where its postcode still "
        "matches the current DE record, otherwise the postcodes.io postcode centroid; geo_source says which.",
        "phone comes from the 2016 OpenDataNI School Locations dataset - the most recent openly licensed source "
        "DE publishes - so it may be out of date.",
        "capacity is the DE 'approved enrolment number' for 2025/26; pupils is the census enrolment. Counts under "
        "five are suppressed by DE and appear as null.",
        "DE publishes no per-school religious character, website, head teacher, or post-primary sex breakdown, so "
        "religious_ethos (except Catholic Maintained), website, head and post-primary gender are null.",
        "Independent (non grant-aided) schools and pre-school playgroups are not in this register; DE publishes "
        "them separately or not at all.",
    ]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--refresh", action="store_true", help="re-download everything instead of using the cache")
    parser.add_argument("--offline", action="store_true", help="fail rather than download anything new")
    args = parser.parse_args()

    generated = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    schools, stats = build_schools(args.refresh, args.offline)

    by_council: dict[str, list[dict]] = {}
    for school in schools:
        by_council.setdefault(school["la"], []).append(school)

    log("")
    entries = []
    for code, name in sorted(COUNCILS.values()):
        rows = sorted(by_council.get(code, []), key=lambda s: (s["name"] or "").lower())
        coords = [(s["lat"], s["lon"]) for s in rows if s["lat"] is not None and s["lon"] is not None]
        bbox = centroid = None
        if coords:
            lats = [c[0] for c in coords]
            lons = [c[1] for c in coords]
            bbox = [
                round(min(lons) - 0.01, 5),
                round(min(lats) - 0.01, 5),
                round(max(lons) + 0.01, 5),
                round(max(lats) + 0.01, 5),
            ]
            centroid = [round(sum(lats) / len(lats), 5), round(sum(lons) / len(lons), 5)]
        # every school in this register is grant-aided except fee-paying preparatory departments
        state = sum(1 for s in rows if "preparatory department" not in (s["type"] or ""))
        size = write_json(
            OUT / "la" / code / "schools.json",
            {
                "generated": generated,
                "sources": sources(),
                "la": {"code": code, "name": name, "nation": "Northern Ireland"},
                "count": len(rows),
                "notes": notes(len(rows)),
                "schools": rows,
            },
        )
        entries.append(
            {
                "la_code": code,
                "name": name,
                "region": "Northern Ireland",
                "nation": "Northern Ireland",
                "lad_codes": [code],
                "school_count": len(rows),
                "state_school_count": state,
                "bbox": bbox,
                "centroid": centroid,
            }
        )
        log(f"  {code}  {name:<38} {len(rows):>4} schools ({state} grant-aided)  {size/1024:6.1f} KB")

    index_size = write_json(
        OUT / "uk" / "northern_ireland_las.json",
        {
            "generated": generated,
            "sources": sources(),
            "notes": [
                f"The 11 Northern Ireland councils, with the grant-aided schools DE recorded at the {CENSUS_DATE} "
                "school census.",
                "la_code and lad_codes are the ONS 2014 local government district code; Northern Ireland councils "
                "are not education authorities - the Education Authority is region-wide.",
                "state_school_count excludes fee-paying preparatory departments of voluntary grammar schools.",
                "bbox is [west, south, east, north] padded by 0.01 degrees; centroid is the mean of school "
                "coordinates.",
            ],
            "las": entries,
        },
    )
    log("")
    log(f"  uk/northern_ireland_las.json  {index_size/1024:.1f} KB")
    log(f"  {len(schools)} schools, {stats['locate_points']} located from the 2016 building points, "
        f"{stats['postcode_centroids']} from postcode centroids, {len(stats['unplaced'])} unplaced")
    for ref, name, postcode in stats["unplaced"]:
        log(f"    unplaced: {ref} {name} ({postcode})")
    if stats["council_mismatch"]:
        log(f"  {len(stats['council_mismatch'])} schools whose postcode centroid falls in a different council "
            "than DE records (kept under DE's council):")
        for ref, name, postcode, de_name, pc_name in stats["council_mismatch"]:
            log(f"    {ref} {name} ({postcode}): DE {de_name} / postcodes.io {pc_name}")
    for ref, council in stats["unknown_council"]:
        log(f"    DROPPED {ref}: unmapped council {council!r}")
    for ref in stats["non_numeric_ref"]:
        log(f"    DROPPED {ref}: non-numeric DE reference")


if __name__ == "__main__":
    main()
